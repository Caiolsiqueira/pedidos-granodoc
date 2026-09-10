import io
import os
import csv
from typing import List, Dict, Any, Tuple
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

# Dados de exemplo realistas para o restaurante
EXEMPLOS_INSUMOS = [
    ("Farinha de Trigo Especial 00", "Farinhas & Massas", "kg", "Pizza"),
    ("Queijo Mussarela Especial Fatiado", "Queijos & Laticínios", "kg", "Pizza, Cozinha"),
    ("Molho de Tomate Pelati Italiano", "Molhos & Bases", "lata", "Pizza"),
    ("Filé Mignon Bovino Limpo Inteiro", "Carnes & Pescados", "kg", "Cozinha"),
    ("Azeite de Oliva Extra Virgem 500ml", "Azeites & Óleos", "garrafa", "Pizza, Cozinha"),
    ("Gin Tanqueray London Dry 750ml", "Destilados & Bebidas", "garrafa", "Bar"),
    ("Gelo Cristalino em Cubos", "Gelo & Bases", "pct", "Bar, Salão"),
    ("Vinho Tinto Malbec Reserva 750ml", "Vinhos Tintos", "garrafa", "Salão"),
    ("Taças de Cristal Degustação 450ml", "Acessórios & Salão", "un", "Salão"),
    ("Manjericão Fresco Orgânico", "Hortifrúti", "mç", "Pizza, Cozinha"),
]

TABELA_UNIDADES = [
    ("kg", "Quilograma", "Farinhas, queijos em barra, carnes, hortifrúti pesado"),
    ("g", "Grama", "Especiarias raras, açafrão, orégano, fermentos"),
    ("litro", "Litro", "Leite, caldos prontos, azeite a granel, xaropes"),
    ("ml", "Mililitro", "Doses, essências, corantes alimentícios concentrados"),
    ("un", "Unidade", "Frutas unitárias (abacaxi, melancia), taças, decanters, utensílios"),
    ("pct", "Pacote", "Massa seca, fermento químico, guardanapos, canudos"),
    ("lata", "Lata", "Tomate pelati, refrigerantes, água tônica, cervejas"),
    ("garrafa", "Garrafa", "Vinhos, destilados, licores, azeites engarrafados"),
    ("cx", "Caixa", "Ovos, morangos frescos, caixas de luvas"),
    ("mç", "Maço", "Hortelã, manjericão, alecrim, tomilho, cebolinha"),
]

TABELA_PRACAS = [
    ("Pizza", "Pizzas, massas, calzones e insumos do forno a lenha"),
    ("Cozinha", "Carnes, risotos, molhos quentes, frituras e pré-preparo"),
    ("Bar", "Coquetelaria, destilados, xaropes, frutas e não alcoólicos"),
    ("Salão", "Adega, vinhos, espumantes, taças, insumos de atendimento"),
    ("Multi-praça", "Pode colocar mais de uma separada por vírgula (ex: 'Pizza, Cozinha' ou 'Bar, Salão')"),
]


def generate_excel_bytes() -> bytes:
    """Gera uma planilha Excel .xlsx profissional e estilizada com abas de insumos e instruções."""
    wb = openpyxl.Workbook()

    # --- ABA 1: INSUMOS (Para preenchimento) ---
    ws1 = wb.active
    ws1.title = "Insumos"
    ws1.views.sheetView[0].showGridLines = True

    # Estilos
    header_fill = PatternFill(start_color="1E293B", end_color="1E293B", fill_type="solid")
    header_font = Font(name="Segoe UI", size=11, bold=True, color="FFFFFF")
    data_font = Font(name="Segoe UI", size=10, color="0F172A")
    border_thin = Border(
        left=Side(style="thin", color="CBD5E1"),
        right=Side(style="thin", color="CBD5E1"),
        top=Side(style="thin", color="CBD5E1"),
        bottom=Side(style="thin", color="CBD5E1"),
    )
    align_center = Alignment(horizontal="center", vertical="center")
    align_left = Alignment(horizontal="left", vertical="center")

    # Cabeçalho da tabela
    headers = [
        "Nome do Produto *",
        "Categoria",
        "Unidade de Medida *",
        "Praça(s) (Opcional)",
    ]
    ws1.append(headers)

    for col_idx in range(1, 5):
        cell = ws1.cell(row=1, column=col_idx)
        cell.fill = header_fill
        cell.font = header_font
        cell.alignment = align_center
        cell.border = border_thin

    # Linhas de exemplo
    for item in EXEMPLOS_INSUMOS:
        ws1.append(list(item))
        row_idx = ws1.max_row
        for col_idx in range(1, 5):
            cell = ws1.cell(row=row_idx, column=col_idx)
            cell.font = data_font
            cell.border = border_thin
            if col_idx in (3, 4):
                cell.alignment = align_center
            else:
                cell.alignment = align_left

    # Ajusta larguras
    ws1.column_dimensions["A"].width = 38
    ws1.column_dimensions["B"].width = 28
    ws1.column_dimensions["C"].width = 22
    ws1.column_dimensions["D"].width = 26

    # --- ABA 2: INSTRUÇÕES E MÉTRICAS ---
    ws2 = wb.create_sheet(title="Instruções & Unidades")
    ws2.views.sheetView[0].showGridLines = True

    # Título do Guia
    ws2.merge_cells("A1:D1")
    title_cell = ws2["A1"]
    title_cell.value = "GUIA DE PREENCHIMENTO - CADASTRO EM MASSA (PEDIDOS GRANODOC)"
    title_cell.font = Font(name="Segoe UI", size=13, bold=True, color="FFFFFF")
    title_cell.fill = PatternFill(start_color="D97706", end_color="D97706", fill_type="solid")
    title_cell.alignment = align_center

    # Explicação das Colunas
    ws2["A3"] = "ORDEM DAS COLUNAS NA PLANILHA:"
    ws2["A3"].font = Font(name="Segoe UI", size=11, bold=True, color="1E293B")

    col_docs = [
        ("1ª Coluna", "Nome do Produto", "Obrigatório", "Nome completo do insumo ou insumo com tamanho/marca"),
        ("2ª Coluna", "Categoria", "Opcional", "Se vazia, será 'Geral'. Se você criar uma nova, o sistema a cadastra automaticamente"),
        ("3ª Coluna", "Unidade de Medida", "Obrigatório", "Veja a tabela abaixo com as siglas permitidas (ex: kg, litro, un, garrafa)"),
        ("4ª Coluna", "Praça(s)", "Opcional", "Pizza, Cozinha, Bar ou Salão. Se omitido, usará a praça selecionada no sistema. Pode colocar mais de uma separada por vírgula"),
    ]
    ws2.append(["Coluna", "Campo", "Status", "Descrição e Regras"])
    header_row_2 = ws2.max_row
    for col_idx in range(1, 5):
        cell = ws2.cell(row=header_row_2, column=col_idx)
        cell.fill = PatternFill(start_color="334155", end_color="334155", fill_type="solid")
        cell.font = Font(name="Segoe UI", size=10, bold=True, color="FFFFFF")
        cell.alignment = align_center
        cell.border = border_thin

    for doc in col_docs:
        ws2.append(list(doc))
        r_idx = ws2.max_row
        for col_idx in range(1, 5):
            cell = ws2.cell(row=r_idx, column=col_idx)
            cell.font = data_font
            cell.border = border_thin
            if col_idx in (1, 3):
                cell.alignment = align_center
            else:
                cell.alignment = align_left

    # Tabela de Unidades Aceitas
    ws2.cell(row=ws2.max_row + 2, column=1, value="UNIDADES DE MEDIDA / MÉTRICAS ACEITAS:").font = Font(name="Segoe UI", size=11, bold=True, color="1E293B")
    ws2.append(["Sigla / Unidade", "Nome Completo", "Exemplos Típicos de Insumos"])
    h_unit_row = ws2.max_row
    for col_idx in range(1, 4):
        cell = ws2.cell(row=h_unit_row, column=col_idx)
        cell.fill = PatternFill(start_color="0284C7", end_color="0284C7", fill_type="solid")
        cell.font = Font(name="Segoe UI", size=10, bold=True, color="FFFFFF")
        cell.alignment = align_center
        cell.border = border_thin

    for u in TABELA_UNIDADES:
        ws2.append(list(u))
        r_idx = ws2.max_row
        for col_idx in range(1, 4):
            cell = ws2.cell(row=r_idx, column=col_idx)
            cell.font = data_font
            cell.border = border_thin
            if col_idx == 1:
                cell.alignment = align_center
                cell.font = Font(name="Segoe UI", size=10, bold=True, color="0369A1")
            else:
                cell.alignment = align_left

    # Tabela de Praças
    ws2.cell(row=ws2.max_row + 2, column=1, value="PRAÇAS OPERACIONAIS VÁLIDAS:").font = Font(name="Segoe UI", size=11, bold=True, color="1E293B")
    ws2.append(["Nome da Praça", "Descrição e Aplicação"])
    h_sec_row = ws2.max_row
    for col_idx in range(1, 3):
        cell = ws2.cell(row=h_sec_row, column=col_idx)
        cell.fill = PatternFill(start_color="059669", end_color="059669", fill_type="solid")
        cell.font = Font(name="Segoe UI", size=10, bold=True, color="FFFFFF")
        cell.alignment = align_center
        cell.border = border_thin

    for s in TABELA_PRACAS:
        ws2.append(list(s))
        r_idx = ws2.max_row
        for col_idx in range(1, 3):
            cell = ws2.cell(row=r_idx, column=col_idx)
            cell.font = data_font
            cell.border = border_thin
            if col_idx == 1:
                cell.alignment = align_center
                cell.font = Font(name="Segoe UI", size=10, bold=True, color="047857")
            else:
                cell.alignment = align_left

    ws2.column_dimensions["A"].width = 20
    ws2.column_dimensions["B"].width = 28
    ws2.column_dimensions["C"].width = 22
    ws2.column_dimensions["D"].width = 50

    output = io.BytesIO()
    wb.save(output)
    return output.getvalue()


def generate_csv_bytes() -> bytes:
    """Gera um arquivo CSV (.csv) com BOM UTF-8 delimitado por ponto-e-vírgula para abertura direta no Excel."""
    output = io.StringIO()
    # Ponto-e-vírgula é o padrão do Excel em língua portuguesa
    writer = csv.writer(output, delimiter=";", quoting=csv.QUOTE_MINIMAL)
    writer.writerow(["Nome do Produto", "Categoria", "Unidade de Medida", "Praça"])
    for item in EXEMPLOS_INSUMOS:
        writer.writerow(list(item))

    # Adiciona BOM UTF-8 (\xef\xbb\xbf) para que o Excel abra sem problemas de acentuação
    csv_str = output.getvalue()
    return ("\ufeff" + csv_str).encode("utf-8")


def parse_uploaded_file(
    file_bytes: bytes,
    filename: str,
    default_setor: str = "Pizza",
    default_unidade: str = "un",
    default_categoria: str = "Geral"
) -> List[Dict[str, Any]]:
    """
    Lê e converte um arquivo .xlsx ou .csv em uma lista de produtos estruturados.
    Reconhece automaticamente as colunas de Nome, Categoria, Unidade e Praça.
    """
    items = []
    fname_lower = filename.lower()

    if fname_lower.endswith((".xlsx", ".xlsm")):
        wb = openpyxl.load_workbook(io.BytesIO(file_bytes), data_only=True)
        # Seleciona a primeira aba ou a aba 'Insumos'
        ws = wb["Insumos"] if "Insumos" in wb.sheetnames else wb.active
        rows = list(ws.iter_rows(values_only=True))
        if not rows:
            return []

        # Tenta mapear o cabeçalho na primeira linha
        header = [str(c).strip().lower() if c is not None else "" for c in rows[0]]
        col_nome = 0
        col_cat = 1
        col_unid = 2
        col_setor = 3

        has_header = False
        for idx, h in enumerate(header):
            if any(k in h for k in ["nome", "produto", "insumo", "item"]):
                col_nome = idx
                has_header = True
            elif any(k in h for k in ["cat", "grupo", "seção"]):
                col_cat = idx
                has_header = True
            elif any(k in h for k in ["unid", "medida", "métrica", "sigla"]):
                col_unid = idx
                has_header = True
            elif any(k in h for k in ["praça", "praca", "setor", "área"]):
                col_setor = idx
                has_header = True

        data_rows = rows[1:] if has_header else rows
        for row in data_rows:
            if not row or not any(row):
                continue
            nome = str(row[col_nome]).strip() if col_nome < len(row) and row[col_nome] is not None else ""
            if not nome or nome.lower() in ["nome do produto", "produto", "nome"]:
                continue

            categoria = (
                str(row[col_cat]).strip()
                if col_cat < len(row) and row[col_cat] is not None and str(row[col_cat]).strip()
                else default_categoria
            )
            unidade = (
                str(row[col_unid]).strip()
                if col_unid < len(row) and row[col_unid] is not None and str(row[col_unid]).strip()
                else default_unidade
            )
            setor_raw = (
                str(row[col_setor]).strip()
                if col_setor < len(row) and row[col_setor] is not None and str(row[col_setor]).strip()
                else default_setor
            )

            items.append({
                "nome": nome,
                "categoria": categoria or default_categoria,
                "unidade": unidade or default_unidade,
                "setor": setor_raw or default_setor,
            })

    elif fname_lower.endswith((".csv", ".txt")):
        # Decodifica arquivo CSV com detecção de encoding (UTF-8 com ou sem BOM, ou ISO-8859-1)
        content = ""
        for enc in ["utf-8-sig", "utf-8", "latin-1", "cp1252"]:
            try:
                content = file_bytes.decode(enc)
                break
            except Exception:
                continue

        lines = content.strip().splitlines()
        if not lines:
            return []

        # Identifica delimitador
        first_line = lines[0]
        delimiter = ";" if ";" in first_line else ("\t" if "\t" in first_line else ",")
        reader = csv.reader(lines, delimiter=delimiter)
        raw_rows = list(reader)

        if not raw_rows:
            return []

        header = [str(c).strip().lower() for c in raw_rows[0]]
        col_nome = 0
        col_cat = 1
        col_unid = 2
        col_setor = 3

        has_header = False
        for idx, h in enumerate(header):
            if any(k in h for k in ["nome", "produto", "insumo", "item"]):
                col_nome = idx
                has_header = True
            elif any(k in h for k in ["cat", "grupo", "seção"]):
                col_cat = idx
                has_header = True
            elif any(k in h for k in ["unid", "medida", "métrica", "sigla"]):
                col_unid = idx
                has_header = True
            elif any(k in h for k in ["praça", "praca", "setor", "área"]):
                col_setor = idx
                has_header = True

        data_rows = raw_rows[1:] if has_header else raw_rows
        for row in data_rows:
            if not row or not any(row):
                continue
            nome = str(row[col_nome]).strip() if col_nome < len(row) else ""
            if not nome or nome.lower() in ["nome do produto", "produto", "nome"]:
                continue

            categoria = str(row[col_cat]).strip() if col_cat < len(row) and str(row[col_cat]).strip() else default_categoria
            unidade = str(row[col_unid]).strip() if col_unid < len(row) and str(row[col_unid]).strip() else default_unidade
            setor_raw = str(row[col_setor]).strip() if col_setor < len(row) and str(row[col_setor]).strip() else default_setor

            items.append({
                "nome": nome,
                "categoria": categoria or default_categoria,
                "unidade": unidade or default_unidade,
                "setor": setor_raw or default_setor,
            })

    return items


def save_template_files(target_dir: str):
    """Salva os arquivos físicos .xlsx e .csv no diretório do projeto para acesso direto."""
    os.makedirs(target_dir, exist_ok=True)
    xlsx_path = os.path.join(target_dir, "modelo_cadastro_insumos.xlsx")
    csv_path = os.path.join(target_dir, "modelo_cadastro_insumos.csv")

    with open(xlsx_path, "wb") as f:
        f.write(generate_excel_bytes())

    with open(csv_path, "wb") as f:
        f.write(generate_csv_bytes())

    return xlsx_path, csv_path


if __name__ == "__main__":
    current_dir = os.path.dirname(os.path.abspath(__file__))
    xlsx, csv_f = save_template_files(current_dir)
    print(f"Modelos gerados com sucesso:\n- {xlsx}\n- {csv_f}")
